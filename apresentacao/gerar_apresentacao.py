"""Gera PPTX editável com notas e prévias/PDF a partir dos mesmos objetos.

Não usa mockups. Capturas vêm de ferramentas/verificar_web.py.
O renderizador Pillow abaixo cobre as formas usadas aqui; não é o motor Office.
"""
import io
import json
import re
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent.parent))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont
from compiler_arena.compiler import compile_source, export_report
from compiler_arena.bytecode import NAMES
from compiler_arena.arena import Arena

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'apresentacao'
BG='101B2D'; WHITE='E8EFF7'; MUTED='ACBCD0'; GREEN='48CFAD'; BLUE='569CF0'; PANEL='192942'
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
FONT='DejaVu Sans'; MONO='DejaVu Sans Mono'


def rect(slide,x,y,w,h,color=PANEL):
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(color); shape.line.fill.background()
    return shape


def text(slide,content,x,y,w=12,size=22,color=WHITE,font=FONT,bold=False):
    # Uma caixa por linha torna geometria e revisão explícitas e editáveis.
    for i,line in enumerate(content.split('\n')):
        box=slide.shapes.add_textbox(Inches(x),Inches(y+i*size/72*1.35),Inches(w),Inches(size/72*1.35))
        tf=box.text_frame; tf.clear(); tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        p=tf.paragraphs[0]; p.space_before=p.space_after=Pt(0)
        run=p.add_run(); run.text=line; run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor.from_string(color)


def slide(title,notes):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    rect(s,.55,.42,.12,.45,GREEN); text(s,title,.85,.38,11.8,30,bold=True)
    text(s,'COMPILER ARENA  /  COMPILADORES',.65,7.08,10,10,MUTED)
    text(s,f'{len(prs.slides):02}',12.1,7.03,.6,13,GREEN)
    s.notes_slide.notes_text_frame.text=notes
    return s


def bullets(s,items,x=.8,y=1.6,w=11.7,size=23,gap=.75):
    for i,item in enumerate(items): text(s,'• '+item,x,y+i*gap,w,size)


def code(s,content,x=.8,y=1.6,w=11.7,size=17):
    lines=content.split('\n'); height=len(lines)*size/72*1.35+.4
    rect(s,x-.15,y-.15,w+.25,height)
    text(s,content,x,y,w,size,font=MONO)


def ir_line(q):
    def f(a):
        if a is None: return ''
        if a['kind']=='constant': return repr(a['value'])
        return {'symbol':'s','temporary':'t','label':'L'}[a['kind']]+str(a['value'])
    rhs=q['op']+' '+f(q['arg1'])+(', '+f(q['arg2']) if q['arg2'] else '')
    return (f(q['result'])+' = ' if q['result'] else '')+rhs


def main():
    OUT.mkdir(exist_ok=True)
    alpha=compile_source((ROOT/'exemplos/alpha.robo').read_text(encoding='utf-8'))
    beta=compile_source((ROOT/'exemplos/beta.robo').read_text(encoding='utf-8'))
    opt=compile_source((ROOT/'exemplos/otimizacao.robo').read_text(encoding='utf-8'))
    export_report(alpha,ROOT/'artefatos/alpha'); export_report(opt,ROOT/'artefatos/otimizacao')
    battle=Arena([alpha['bytecode'],beta['bytecode']],max_ticks=300); final=battle.run()
    tests=(ROOT/'artefatos/testes.txt').read_text(encoding='utf-8') if (ROOT/'artefatos/testes.txt').exists() else 'Testes não executados nesta geração'
    match=re.search(r'Ran (\d+) tests in ([\d.]+)s',tests)
    test_result=f'{match[1]} testes / {match[2]} s / '+('OK' if '\nOK' in tests else 'ver relatório') if match else 'Sem resultado registrado'
    s=slide('Da linguagem-fonte à batalha', 'Apresente os integrantes reais preenchendo os campos. Explique que a batalha é a saída visível de uma cadeia real de compilação, que será percorrida com o robô Alpha.')
    text(s,'Compiler Arena',.85,1.7,11.5,48,GREEN,bold=True)
    text(s,'Uma linguagem. Um compilador. Robôs em execução.',.9,2.65,11.5,25)
    text(s,'Integrantes: [preencher]\nProfessor(a): [preencher]\nInstituição: [preencher]',.9,4.3,11.5,21,MUTED)

    s=slide('Objetivo e requisitos', 'Diferencie o PDF das extensões: o PDF pede as etapas reais, API, arena e ticks; o pedido detalha tipos, escopos, três otimizações, GUI e slides. O PDF permite variações de sintaxe. Não executar código Python do usuário.')
    bullets(s,['PDF: compilar, validar e executar uma linguagem de robôs.', 'Semântica, IR, otimização e objeto são o centro do trabalho.', 'API mínima + arena funcional + execução justa por ticks.', 'Extensões: tipos, escopos, JSON, GUI, testes e relatórios.', 'Sem exec(), eval() ou execução direta da AST.'])

    s=slide('Arquitetura: a transformação é real', 'Siga cada seta: lexer gera tokens; parser gera AST; semântica anota; Builder gera quádruplas; otimizador cria nova IR; gerador resolve registradores e rótulos; VM emite comandos; arena resolve efeitos. Todos os quadros são editáveis.')
    labels=['Fonte UTF-8','Tokens','AST','Semântica','IR tipada','Otimização','Bytecode','VM / arena']
    for i,label in enumerate(labels):
        x=.8+(i%4)*3.13; y=1.8+(i//4)*1.8
        rect(s,x,y,2.75,.85); text(s,label,x+.15,y+.25,2.5,21,GREEN if i in (0,7) else WHITE)
        if i%4!=3: text(s,'>',x+2.85,y+.25,.3,20,MUTED)
    text(s,'CLI e núcleo independem da interface gráfica.',.85,6,11,22,MUTED)

    s=slide('RoboLang: o robô Alpha', 'Este é o fonte real de exemplos/alpha.robo. O mesmo robô fornece os tokens, AST, símbolo, IR, bytecode e batalha mostrados depois. O comentário foi omitido só para caber. Scan cede até o próximo tick e a direção pode envelhecer.')
    source='\n'.join(line for line in alpha['source'].splitlines() if not line.strip().startswith('//'))
    code(s,source,.85,1.45,8.2,15)
    text(s,'Um robô por arquivo\nDeclaração antes do uso\nParênteses e chaves\nComandos sem retorno',9.45,2,3,16,MUTED)
    text(s,'SHA '+alpha['sha256'][:12],9.45,5.5,3,14,GREEN,font=MONO)

    s=slide('Tipos, operadores e controle', 'Explique int→float e por que bool não é int em RoboLang. A divisão produz float, então x/=2 para um int é inválido. Toda declaração é inicializada de novo quando o fluxo passa por ela. Break e continue atuam no laço mais interno.')
    bullets(s,['int, float finito, bool e string Unicode.', 'Única promoção implícita: int para float.', 'Precedência: unários, produto, soma, ordem, igualdade, &&, ||.', 'if / else if / else, while, break e continue.', 'Constantes, blocos léxicos e operadores compostos.'],size=22)

    s=slide('Análise léxica: caracteres → tokens', 'Mostre linha e coluna e a distinção entre lexema e valor. O lexer foi ampliado a partir do original, preservando cursor, máquina de estados e tokens. Não usa regex ou geradores. Floats, comentários e escapes estritos foram acrescentados.')
    tokens=alpha['tokens']; first=next(i for i,t in enumerate(tokens) if t['lexeme']=='int')
    lines=['TIPO               LEXEMA       VALOR     POSIÇÃO']
    for t in tokens[first:first+7]: lines.append(f'{t["type"]:18} {t["lexeme"]:12} {str(t["value"]):9} {t["line"]}:{t["column"]}')
    code(s,'\n'.join(lines),size=17)
    text(s,'Scanner manual • comentários preservam posições • maior operador primeiro',.85,5.9,12,18,MUTED)

    s=slide('Parser recursivo e AST', 'A árvore é gerada pelo parser real. Cada função representa um nível de precedência ou regra. O nó DECL é uma extensão da gramática original. Há recuperação em ; e }, teto de 50 diagnósticos e limite de profundidade de 100.')
    decl=alpha['ast']['filhos'][0]['filhos'][0]
    lines=[]
    def ast(n,depth=0):
        lines.append('  '*depth+n['tipo']+(f' [{n["valor"]}]' if n['valor'] is not None else '')+f' @{n["line"]}:{n["column"]}')
        for c in n['filhos']: ast(c,depth+1)
    ast(decl)
    code(s,'\n'.join(lines),size=22)
    text(s,'A forma de 10 + 20 desaparece em etapas posteriores.',.85,5.5,12,22,MUTED)

    s=slide('Semântica: nomes, tipos e escopos', 'O símbolo vem do relatório de Alpha. No escopo interno de um exemplo sombreado, o mesmo nome ganha outro ID. A inicialização resolve nomes antes da inserção do símbolo novo. A AST recebe meta.type e meta.symbol, consumidos pelo gerador de IR.')
    sym=alpha['symbols'][0]
    code(s,'\n'.join(f'{key}: {value}' for key,value in sym.items()),.85,1.5,5.5,18)
    text(s,'Tabela de símbolos real\n\nCada declaração tem um ID.\nO backend usa esse ID.\nTipos são validados antes da IR.',7.2,1.8,5.5,20)

    s=slide('Diagnóstico e bloqueio do inválido', 'Este diagnóstico é produzido ao compilar o mesmo Alpha com fire(true). Nenhum bytecode válido é retornado. A GUI invalida resultados quando o texto muda e descarta respostas de revisões anteriores. O programa inválido não pode ser adicionado à batalha.')
    invalid=compile_source(alpha['source'].replace('fire(3)','fire(true)'))
    code(s,'fire(true);',.85,1.6,11.7,26)
    message=invalid['diagnostics'][0]
    # quebra manual para legibilidade
    text(s,message.split(': ',1)[0]+':\n'+message.split(': ',1)[1],.85,3.2,11.7,19,'F4A1AC')
    text(s,'Semântica: erro  |  IR: não executada  |  Bytecode: ausente',.85,5.55,11.7,22,GREEN)

    s=slide('IR tipada: quádruplas e saltos', 'Os quatro primeiros registros são da inicialização e do movimento de Alpha. O rótulo seguinte inicia o while. Cada operando tem categoria explícita, tipo e valor; cada Quad tem origem. JZ e JMP codificam a decisão e o retorno, em vez de interpretar a AST.')
    code(s,'\n'.join(ir_line(q) for q in alpha['ir'][:9]),.85,1.5,11.7,19)
    text(s,'op / arg1 / arg2 / result + tipo + linha/coluna',.85,5.7,11.7,22,MUTED)

    s=slide('IR original e otimizada: o mesmo Alpha', 'Esta comparação usa o prefixo real da IR de Alpha. As demais instruções ainda representam laço e combate. A contagem total inclui rótulos IR, portanto não deve ser confundida com instruções executáveis ou ganho de tempo. O exemplo isolado otimizacao.robo reduz 5 para 3.')
    text(s,'Original',.85,1.5,5.5,24,BLUE); text(s,'Otimizada',7,1.5,5.5,24,GREEN)
    code(s,'\n'.join(ir_line(q) for q in alpha['ir'][:4]),.85,2.2,5.3,17)
    code(s,'\n'.join(ir_line(q) for q in alpha['optimized_ir'][:2]),7,2.2,5.3,17)
    m=alpha['metrics']; om=opt['metrics']
    text(s,f'Alpha completo: {m["before"]} → {m["after"]} instruções IR ({m["reduction_percent"]:.2f}%).',.85,4.6,11.7,23)
    text(s,f'Trecho isolado: {om["before"]} → {om["after"]} ({om["reduction_percent"]:.0f}%). Não é medida de tempo.',.85,5.45,11.7,21,MUTED)

    s=slide('Otimização: três técnicas, com limites', 'Mostre o log de regras aplicado em artefatos/alpha/optimizations.json. CFG identifica líderes e sucessores. Não atravessamos arestas com o ambiente de constantes. O orçamento de instruções explica por que menos instruções pode mudar a cronologia de uma batalha.')
    bullets(s,['Constant folding: 10 + 20 torna-se COPY 30.', 'Simplificação algébrica: int x + 0 torna-se COPY x.', 'Propagação de constantes dentro de blocos básicos.', 'Preserva sensores, comandos, tipos e erros observáveis.', 'Equivalência funcional não garante equivalência temporal.'],size=22)

    s=slide('Bytecode próprio e máquina virtual', 'Este código é o bytecode real otimizado de Alpha. Os índices substituem rótulos e as leituras apontam constantes ou registradores. A VM valida JSON, mantém PC e memória própria e não importa o compilador. Source permite associar erro de execução ao fonte.')
    rows=alpha['bytecode']['instructions']
    code(s,'\n'.join(f'{i:02} {NAMES[r["opcode"]]:10} {r["args"]}' for i,r in enumerate(rows[:9])),.85,1.5,11.7,17)
    text(s,'JSON validado • registradores isolados • PC persistente • sem AST',.85,5.7,11.7,20,MUTED)

    s=slide('Ticks: orçamento e ações simultâneas', 'Cada opcode consome orçamento, inclusive saltos e HALT. A emissão de um comando cede a VM. Scan é um comando, e seu resultado só entra no snapshot seguinte. Todos os ataques usam o início do tick; custos e dano são acumulados antes da eliminação.')
    bullets(s,['1  Snapshot: posições, energia e último scan.', '2  VMs: orçamento por robô; até um comando por tick.', '3  Ações: movimento, rotação, scan e tiro.', '4  Custos e dano simultâneos; eliminação.', '5  Novo estado; scan publicado para o próximo tick.'],size=22)
    text(s,'Loop infinito sem comando pausa ao esgotar o orçamento.',.85,6.05,11.7,20,GREEN)

    s=slide('Editor: etapas reais e fonte identificada', 'Captura real da interface HTML no Chromium headless. A esquerda está Alpha; a direita, a IR real. O navegador envia fonte para a API Python e descarta respostas de revisões anteriores. Não é um mockup. Amplie o aplicativo durante a demonstração ao vivo.')
    capture=OUT/'capturas/editor.png'
    if not capture.exists(): raise FileNotFoundError('Gere capturas reais com ferramentas/verificar_web.py')
    s.shapes.add_picture(str(capture),Inches(.9),Inches(1.4),width=Inches(11.5))

    s=slide('Arena: uma única fonte de verdade', 'Captura real da arena HTML após quatro ticks da mesma batalha Alpha/Beta. A única simulação vive no servidor Python; o canvas interpola snapshots e anima disparos com o alvo e o dano informados pelo servidor. A velocidade inicial é 1 tick por segundo. O renderizador inverte y para desenhar. A simulação usa coordenadas cartesianas. Mostre na demonstração os botões pausar e passo. Uma VM finalizada pode ser atingida fisicamente.')
    capture=OUT/'capturas/arena.png'
    if not capture.exists(): raise FileNotFoundError('Gere a captura real da arena')
    s.shapes.add_picture(str(capture),Inches(2.45),Inches(1.3),width=Inches(8.4))
    text(s,f'Execução de referência: Alpha venceu no tick {final["tick"]}.',.85,6.65,11.7,19,GREEN)

    s=slide('Testes e demonstração integrada', 'Resultado extraído do relatório unittest real, não de uma estimativa de cobertura. O teste gráfico verifica HTTP, descarte de resultado antigo, um passo e reinício. Abra a GUI e siga docs/defesa_demonstracao.md. O tempo reportado é da suíte, não desempenho das otimizações.')
    text(s,test_result,.85,1.7,11.7,32,GREEN,bold=True)
    bullets(s,['Léxico → parser → semântica → IR → VM → arena.', 'Curto-circuito, escopos, erros, orçamento e simultaneidade.', 'Teste gráfico: editor, revisão do texto, HTTP, passo e reinício.', 'Ao vivo: erro semântico → correção → IR → bytecode → batalha.'],y=2.8,size=21,gap=.8)

    s=slide('Limitações e evoluções', 'Encerre com os limites reais: não há propagação global, nem funções, nem física contínua além de trajetórias lineares por tick. Explicite a validação gráfica em Linux, sem afirmar teste Windows. Evoluções planejadas não estão implementadas. Use as perguntas e respostas da documentação para preparar a equipe.')
    text(s,'Implementado',.85,1.6,5.6,25,GREEN)
    text(s,'Cadeia completa e inspecionável\nCLI + editor + arena\nTestes e exportação JSON\nSlides gerados dos artefatos',.85,2.4,5.6,20)
    text(s,'Limites / próximos passos',7,1.6,5.6,25,BLUE)
    text(s,'Sem funções ou propagação global\nScan pode ficar desatualizado\nTiros instantâneos, sem projéteis\nTeste web local: Chromium / Linux',7,2.4,5.6,20)
    text(s,'A defesa deve explicar as transformações, não só a batalha.',.85,5.9,11.7,22,MUTED)

    path=OUT/'apresentacao_compiler_arena.pptx'; prs.save(path)
    render(path)
    (OUT/'dados_slides.json').write_text(json.dumps({'sha256':alpha['sha256'],'metrics':m,'isolated_metrics':om,'tests':test_result,'battle':final},ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'{path}: {len(prs.slides)} slides editáveis com notas; prévias e PDF gerados.')


def render(path):
    """Renderiza somente os objetos usados nesta apresentação e checa limites."""
    doc=Presentation(path); scale=120; out=OUT/'previas'; out.mkdir(exist_ok=True)
    W,H=1600,900; pages=[]; issues=[]
    def px(emu): return round(emu/914400*scale)
    def font(name,size,bold):
        family='DejaVuSansMono' if 'Mono' in name else 'DejaVuSans'
        filename=family+('-Bold' if bold else '')+'.ttf'
        try: return ImageFont.truetype(filename,round(size/72*scale))
        except OSError: return ImageFont.truetype('arial.ttf',round(size/72*scale))
    for n,s in enumerate(doc.slides,1):
        canvas=Image.new('RGB',(W,H),'#'+BG); draw=ImageDraw.Draw(canvas)
        text_rects=[]
        for shape in s.shapes:
            x,y,w,h=px(shape.left),px(shape.top),px(shape.width),px(shape.height)
            if x<0 or y<0 or x+w>W+2 or y+h>H+2: issues.append(f'slide {n}: forma fora da tela')
            if shape.shape_type==13:
                picture=Image.open(io.BytesIO(shape.image.blob)).convert('RGB').resize((w,h),Image.Resampling.LANCZOS); canvas.paste(picture,(x,y)); continue
            if not shape.has_text_frame: continue
            try:
                if shape.fill.type==1: draw.rectangle((x,y,x+w,y+h),fill='#'+str(shape.fill.fore_color.rgb))
            except (AttributeError,TypeError): pass
            for para in shape.text_frame.paragraphs:
                if not para.runs: continue
                run=para.runs[0]; fs=run.font.size.pt if run.font.size else 22
                f=font(run.font.name or FONT,fs,run.font.bold); content=para.text
                color='#'+str(run.font.color.rgb)
                bounds=draw.textbbox((x,y),content,font=f,anchor='lt')
                if bounds[2]>x+w+2 or bounds[3]>y+h+2: issues.append(f'slide {n}: texto pode cortar: {content[:50]}')
                if content:
                    for prev in text_rects:
                        if max(bounds[0],prev[0])<min(bounds[2],prev[2]) and max(bounds[1],prev[1])<min(bounds[3],prev[3]): issues.append(f'slide {n}: textos sobrepostos: {content[:30]}')
                    text_rects.append(bounds)
                draw.text((x,y),content,font=f,fill=color,anchor='lt')
        canvas.save(out/f'slide_{n:02}.png'); pages.append(canvas)
    pages[0].save(OUT/'apresentacao_compiler_arena.pdf','PDF',resolution=120,save_all=True,append_images=pages[1:])
    thumbs=Image.new('RGB',(1600,6*310),'#'+BG)
    for i,page in enumerate(pages): thumbs.paste(page.resize((533,300)),((i%3)*533,(i//3)*310))
    thumbs.save(out/'contato.png')
    (OUT/'revisao.txt').write_text('Renderização dos objetos PPTX via Pillow, não pelo motor Office.\nSlides: '+str(len(pages))+'\nChecagens geométricas de textos/formas: '+('sem cortes ou sobreposições detectados' if not issues else '\n'.join(issues))+'\n',encoding='utf-8')
    if issues: raise RuntimeError('Revise os slides: '+str(issues))


if __name__=='__main__': main()
